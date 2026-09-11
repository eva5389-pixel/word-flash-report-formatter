"""Lossless text-first formatting for financial-news email drafts."""
from dataclasses import dataclass
from email.message import EmailMessage
from email import policy
from html import escape
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

GREETING = '各位理專同仁大家好!!以下為本日重要財金新聞匯集'
MAX_BYTES = 10 * 1024 * 1024
DATE = re.compile(r'^\d{4}[/年.-]\d{1,2}[/月.-]\d{1,2}')
NOISE = re.compile(r'^(?:本文共\s*[\d,]+\s*字|\d{1,2}:\d{2})$')
SOURCE = re.compile(r'^(?:來源[：:]|.*(?:日報|新聞網|通訊社|中央社|路透社|Reuters|Bloomberg).*(?:記者|報導|編譯))')

@dataclass
class Article:
    title: str
    paragraphs: list[str]

@dataclass
class Report:
    greeting: str
    articles: list[Article]
    removed: int = 0
    inferred: bool = False


def extract_upload(name: str, data: bytes) -> tuple[str, list[str]]:
    if len(data) > MAX_BYTES:
        raise ValueError('檔案請小於 10 MB。')
    suffix = Path(name).suffix.lower()
    if suffix == '.txt':
        if data.startswith((b'\xff\xfe', b'\xfe\xff')):
            encodings = ['utf-16']
        else:
            encodings = ['utf-8-sig', 'cp950']
        for encoding in encodings:
            try:
                text = data.decode(encoding)
                if '\x00' in text:
                    raise ValueError('文字檔含無法辨識的字元，請另存為 UTF-8 TXT。')
                return text, []
            except UnicodeDecodeError:
                continue
        raise ValueError('無法辨識文字編碼，請另存為 UTF-8 TXT。')
    if suffix != '.docx':
        raise ValueError('請上傳 DOCX 或 TXT 檔案。舊版 DOC 請先另存為 DOCX。')
    try:
        with ZipFile(BytesIO(data)) as package:
            if sum(i.file_size for i in package.infolist()) > 50 * 1024 * 1024:
                raise ValueError('解壓縮後的 Word 檔過大，請分成較小檔案。')
        doc = Document(BytesIO(data))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('無法讀取 Word 檔，請確認檔案未損毀或加密。') from exc
    lines = []
    def walk(parent):
        for block in parent.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    style = block.style.name if block.style else ''
                    if style.startswith(('Heading', 'Title', '標題')):
                        text = '## ' + text
                    lines.append(text)
            elif isinstance(block, Table):
                for row in block.rows:
                    seen = set()
                    for cell in row.cells:
                        if cell._tc not in seen:
                            seen.add(cell._tc)
                            walk(cell)
    walk(doc)
    warnings = []
    if doc.element.xpath('.//w:drawing | .//w:pict | .//w:txbxContent'):
        warnings.append('此版本擷取文字與表格文字；圖片、圖表及文字方塊不會帶入信件，請寄送前補上。')
    return '\n\n'.join(lines), warnings


def parse_report(text: str, greeting: str = GREETING, clean: bool = False) -> Report:
    lines = [line.strip() for line in text.replace('\r\n', '\n').replace('\r', '\n').split('\n') if line.strip()]
    if not lines:
        raise ValueError('請先上傳檔案或貼上新聞內容。')
    if lines[0] == GREETING or (lines[0].startswith('各位') and '大家好' in lines[0]):
        lines.pop(0)
    if not lines:
        raise ValueError('目前只有開場問候，請加入新聞內容。')
    explicit = [i for i, line in enumerate(lines) if line.startswith('## ')]
    starts = list(explicit)
    if not explicit:
        # Recognize the supplied news-copy structure without guessing from length alone.
        last_metadata = -1
        for i, line in enumerate(lines):
            if NOISE.fullmatch(line) and line.startswith('本文共'):
                lower = max(last_metadata + 1, i - 4)
                candidates = [j for j in range(lower, i)
                              if not DATE.match(lines[j]) and not NOISE.fullmatch(lines[j])
                              and not SOURCE.match(lines[j])
                              and len(lines[j]) <= 150
                              and not lines[j].endswith(('。', '！', '？', '!', '?'))
                              and not re.search(r'(?:路透|美聯社|示意圖|提供|攝影)[）)]?$', lines[j])]
                if candidates:
                    starts.append(candidates[0])
                last_metadata = i
    starts = sorted(set(starts))
    inferred = not starts
    if not starts or starts[0] != 0:
        # Keep unmatched leading text rather than dropping it.
        starts.insert(0, 0)
    articles = []
    removed = 0
    for idx, start in enumerate(starts):
        end = starts[idx + 1] if idx + 1 < len(starts) else len(lines)
        heading = lines[start]
        if heading.startswith('## '):
            heading = heading[3:].strip()
        # A prose-only input is a single untitled article; don't lose its first paragraph.
        prose = (inferred and (len(heading) > 150 or heading.endswith(('。', '！', '？')))) or (explicit and start not in explicit)
        body = lines[start:end] if prose else lines[start + 1:end]
        paragraphs = []
        for line in body:
            if clean and (NOISE.fullmatch(line) or (paragraphs and re.sub(r'\s+', '', line) == re.sub(r'\s+', '', paragraphs[-1]))):
                removed += 1
                continue
            paragraphs.append(line)
        articles.append(Article('' if prose else heading, paragraphs))
    return Report(greeting.strip(), articles, removed, inferred)


def plain_email(report: Report) -> str:
    blocks = [report.greeting] if report.greeting else []
    for article in report.articles:
        if article.title:
            blocks.append(article.title)
        blocks.extend(article.paragraphs)
    return '\n\n'.join(blocks) + '\n'


def html_email(report: Report, subject: str, font: str = '微軟正黑體', size: int = 16) -> str:
    fonts = {'微軟正黑體': "'Microsoft JhengHei','PingFang TC',Arial,sans-serif", '標楷體': "DFKai-SB,BiauKai,'KaiTi',serif"}
    family = fonts.get(font, fonts['微軟正黑體'])
    size = max(12, min(24, int(size)))
    pstyle = f'margin:0 0 16px;font-size:{size}px;line-height:1.8;overflow-wrap:anywhere;'
    content = [f'<p style="{pstyle}">{escape(report.greeting)}</p>'] if report.greeting else []
    for article in report.articles:
        if article.title:
            content.append(f'<h2 style="margin:30px 0 16px;font-size:{size + 3}px;line-height:1.6;font-weight:bold;color:#17365d;overflow-wrap:anywhere;">{escape(article.title)}</h2>')
        for paragraph in article.paragraphs:
            meta = bool(DATE.match(paragraph) or SOURCE.match(paragraph) or NOISE.fullmatch(paragraph))
            style = f'margin:0 0 8px;font-size:{max(size - 2, 12)}px;line-height:1.6;color:#595959;overflow-wrap:anywhere;' if meta else pstyle
            content.append(f'<p style="{style}">{escape(paragraph)}</p>')
    return ('<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{escape(subject)}</title></head><body style="margin:0;background:#ffffff;color:#222222;">'
            f'<div style="max-width:800px;margin:0 auto;padding:24px;background:#ffffff;color:#222222;font-family:{family};">'
            + ''.join(content) + '</div></body></html>')


def eml_email(report: Report, subject: str, font: str = '微軟正黑體', size: int = 16) -> bytes:
    message = EmailMessage(policy=policy.SMTP)
    message['Subject'] = ' '.join(subject.splitlines()).strip() or '重要財金新聞匯集'
    message['X-Unsent'] = '1'
    message.set_content(plain_email(report))
    message.add_alternative(html_email(report, subject, font, size), subtype='html')
    return message.as_bytes()
