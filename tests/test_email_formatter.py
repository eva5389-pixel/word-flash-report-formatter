import unittest
from email import policy
from email.parser import BytesParser
from io import BytesIO
from docx import Document
from email_formatter import GREETING, extract_upload, parse_report, plain_email, html_email, eml_email

NEWS = '新聞標題一\n\n本文共100字\n00:00\n2026/09/07 02:48:23\n經濟日報 記者甲／台北報導\n\n完整內容第一段。\n\n新聞標題二\n\n圖片說明。 路透\n圖片說明。 路透\n本文共200字\n00:00\n2026/09/07 02:45:41\n經濟日報 記者乙／台北報導\n\n完整內容第二段。'

class EmailTests(unittest.TestCase):
    def test_news_segmentation_preserves_all_text(self):
        original = GREETING + '\n\n' + NEWS
        report = parse_report(original)
        self.assertEqual([a.title for a in report.articles], ['新聞標題一', '新聞標題二'])
        self.assertEqual(''.join(original.split()), ''.join(plain_email(report).split()))

    def test_optional_cleanup(self):
        result = parse_report(NEWS, clean=True)
        self.assertEqual(result.removed, 5)
        plain = plain_email(result)
        self.assertIn('2026/09/07', plain)
        self.assertIn('記者甲', plain)
        self.assertIn('完整內容第二段。', plain)
        self.assertEqual(plain.count('圖片說明。 路透'), 1)

    def test_explicit_titles_and_prose(self):
        result = parse_report('前言段落。\n## 第一篇\n內文。\n## 第二篇\n內容。')
        self.assertIn('前言段落。', plain_email(result))
        self.assertEqual(result.articles[-1].title, '第二篇')
        result = parse_report('沒有標題但需要保留的內容。')
        self.assertEqual(result.articles[0].title, '')
        self.assertIn('沒有標題但需要保留的內容。', plain_email(result))

    def test_html_escape_and_email_draft(self):
        result = parse_report('## <script>alert(1)</script>\n<img src=x onerror=bad>')
        html = html_email(result, '<test>')
        self.assertNotIn('<script>', html)
        self.assertNotIn('<img', html)
        message = BytesParser(policy=policy.default).parsebytes(eml_email(result, '主旨\r\nBcc: bad@example.com'))
        self.assertIsNone(message['Bcc'])
        self.assertIsNone(message['To'])
        self.assertEqual(message['X-Unsent'], '1')
        self.assertIn('alert(1)', message.get_body(preferencelist=('plain',)).get_content())
        self.assertIn('&lt;script&gt;', message.get_body(preferencelist=('html',)).get_content())

    def test_docx_order_and_chinese_encoding(self):
        doc = Document()
        doc.add_heading('第一篇', 1)
        doc.add_paragraph('內文一。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格資料'
        doc.add_paragraph('內文二。')
        stream = BytesIO(); doc.save(stream)
        text, warnings = extract_upload('news.docx', stream.getvalue())
        self.assertEqual(text, '## 第一篇\n\n內文一。\n\n表格資料\n\n內文二。')
        self.assertEqual(warnings, [])
        for encoding in ['utf-8-sig', 'utf-16', 'cp950']:
            self.assertEqual(extract_upload('news.txt', '財金新聞'.encode(encoding))[0], '財金新聞')

    def test_invalid_upload_and_empty_content(self):
        for name, data in [('bad.docx', b'bad'), ('bad.pdf', b'bad'), ('big.txt', b'a' * (10*1024*1024 + 1))]:
            with self.assertRaises(ValueError): extract_upload(name, data)
        for text in ['', ' \n', GREETING]:
            with self.assertRaises(ValueError): parse_report(text)

if __name__ == '__main__': unittest.main()
